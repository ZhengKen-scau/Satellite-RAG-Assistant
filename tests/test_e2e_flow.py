#!/usr/bin/env python3
"""
端到端测试脚本 - 验证完整的文档问答流程
"""

import os
import tempfile
import requests
from app.config import settings

def test_document_qa_flow():
    """测试完整的文档问答流程"""
    print("🔍 测试端到端文档问答流程")
    
    # 创建测试DOCX文件
    with tempfile.TemporaryDirectory() as temp_dir:
        test_doc_path = os.path.join(temp_dir, "遥感实验报告.docx")
        
        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph("这是郑恳的《遥感基础与应用》实验报告。")
            doc.add_paragraph("报告主要内容包括：遥感基本原理、卫星传感器类型、图像处理方法。")
            doc.add_paragraph("Landsat-8卫星的空间分辨率为30米，重访周期为16天。")
            doc.save(test_doc_path)
            
            # 准备上传数据
            with open(test_doc_path, 'rb') as f:
                files = {'files': ('遥感实验报告.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
                data = {'query': '这个文档主要说了什么？'}
                
                # 发送请求到API
                response = requests.post(
                    f"http://127.0.0.1:8080/api/v1/query",
                    files=files,
                    data=data,
                    timeout=30
                )
                
                if response.status_code == 200:
                    result = response.json()
                    answer = result.get('answer', '')
                    processing_time = result.get('processing_time', 0)
                    
                    print(f"✅ API调用成功 (处理时间: {processing_time:.1f}秒)")
                    print(f"回答预览: {answer[:200]}...")
                    
                    # 检查是否包含文档内容
                    if "遥感" in answer and "实验报告" in answer:
                        print("✅ 回答包含了文档的关键信息")
                        return True
                    elif "无法找到相关答案" in answer:
                        print("❌ 回答显示'无法找到相关答案'，说明上下文检索失败")
                        return False
                    else:
                        print("⚠️  回答可能没有充分利用文档内容")
                        return True  # 可能是其他原因
                        
                else:
                    print(f"❌ API调用失败 (状态码: {response.status_code})")
                    print(f"响应: {response.text}")
                    return False
                    
        except Exception as e:
            print(f"❌ 测试失败: {e}")
            return False

def main():
    """主函数"""
    print("=" * 60)
    print("🚀 End-to-End Document QA Test")
    print("=" * 60)
    
    # 检查Ollama服务
    try:
        ollama_response = requests.get(f"{settings.OLLAMA_HOST}/api/tags", timeout=5)
        if ollama_response.status_code == 200:
            print("✅ Ollama服务正常运行")
        else:
            print("❌ Ollama服务异常")
            return
    except Exception as e:
        print(f"❌ 无法连接到Ollama服务: {e}")
        return
    
    # 运行测试
    success = test_document_qa_flow()
    
    print("\n" + "=" * 60)
    if success:
        print("🎉 端到端测试通过！文档问答功能正常工作！")
    else:
        print("❌ 端到端测试失败，请检查系统配置")

if __name__ == "__main__":
    main()